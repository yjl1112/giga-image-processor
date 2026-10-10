#!/usr/bin/env python3
"""处理 GIGA 素材包图片：读取 Excel 模板，解压素材 ZIP，裁剪填满 1600x1600 并导出。"""

from __future__ import annotations

import argparse
import csv
import os
import re
import zipfile
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from PIL import Image, ImageOps, UnidentifiedImageError

IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff"}
DEFAULT_TEMPLATE = Path("data/input/GIGA素材包下载与图片处理模板.xlsx")
DEFAULT_SHEET = "填写模板"
DEFAULT_INPUT_DIR = Path("data/raw/giga_materials")
DEFAULT_OUTPUT_DIR = Path("data/image_exports")
DEFAULT_REPORT_DIR = Path("data/reports")

YES_VALUES = {"是", "y", "yes", "true", "1"}
NO_VALUES = {"否", "n", "no", "false", "0", ""}
SUPPORTED_PROCESS_METHODS = {
    "裁剪填满1600x1600",
    "裁剪填满 1600x1600",
    "cover_1600",
    "crop_cover_1600",
}

FIELD_ALIASES = {
    "amazon_sku": ["亚马逊SKU（推荐必填）", "Amazon SKU", "amazon_sku"],
    "expected_item_code": ["GIGA Item Code（必填）", "GIGA Item Code", "Item Code", "期望Item Code"],
    "actual_item_code": [
        "页面实际Item Code（可选）",
        "实际Item Code（可选）",
        "素材包实际Item Code（可选）",
        "下载页实际Item Code（可选）",
        "实际Item Code",
    ],
    "product_url": ["GIGA商品链接（可选）", "GIGA商品链接", "商品链接", "giga_product_url"],
    "product_name": ["商品名称（可选）", "商品名称", "product_name"],
    "download_material": ["是否下载素材包（必填：是/否）", "是否下载素材包", "download_material"],
    "process_mode": ["处理方式（必填）", "处理方式", "process_mode"],
    "output_prefix": ["输出文件名前缀（推荐填亚马逊SKU）", "输出文件名前缀", "输出前缀", "output_name"],
    "main_image_only": ["是否只处理主图（必填：是/否）", "是否只处理主图", "main_image_only"],
    "note": ["备注/特殊要求（可选）", "备注", "特殊要求", "notes"],
}


@dataclass
class ProcessResult:
    export_zip: Path
    report_csv: Path
    image_count: int


def clean_cell(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_yes_no(value) -> bool:
    text = clean_cell(value).lower()
    if text in YES_VALUES:
        return True
    if text in NO_VALUES:
        return False
    raise ValueError(f"无法识别是/否值：{value}")


def field_value(row: dict[str, str], field: str) -> str:
    for alias in FIELD_ALIASES[field]:
        value = clean_cell(row.get(alias))
        if value:
            return value
    return ""


def read_xlsx_dict_rows(path: Path, sheet_name: str = DEFAULT_SHEET) -> Iterable[tuple[int, dict[str, str]]]:
    wb = load_workbook(path, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Excel 中找不到 sheet：{sheet_name}；实际 sheet：{', '.join(wb.sheetnames)}")
        ws = wb[sheet_name]
        headers = [clean_cell(cell.value) for cell in ws[1]]
        for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            row = {header: clean_cell(value) for header, value in zip(headers, values) if header}
            if any(row.values()):
                yield row_number, row
    finally:
        wb.close()


def validate_process_mode(value: str, row_number: int) -> None:
    if clean_cell(value) not in SUPPORTED_PROCESS_METHODS:
        supported = "、".join(sorted(SUPPORTED_PROCESS_METHODS))
        raise ValueError(f"第 {row_number} 行处理方式不支持：{value}；当前支持：{supported}")


def match_material_zip(input_dir: Path, actual_item_code: str) -> Path | None:
    code = actual_item_code.lower()
    candidates = [path for path in input_dir.glob("*.zip") if code in path.name.lower()]
    image_file_candidates = [path for path in candidates if "image+file" in path.name.lower()]
    selected = image_file_candidates or candidates
    if not selected:
        return None
    return sorted(selected, key=lambda path: (path.stat().st_mtime, path.name))[-1]


def build_rows_from_template(
    *,
    template_path: Path,
    input_dir: Path,
    sheet_name: str = DEFAULT_SHEET,
) -> list[dict]:
    rows: list[dict] = []
    missing_archives: list[str] = []
    for row_number, row in read_xlsx_dict_rows(template_path, sheet_name):
        download_value = field_value(row, "download_material")
        if not normalize_yes_no(download_value):
            continue

        expected_item_code = field_value(row, "expected_item_code")
        if not expected_item_code:
            raise ValueError(f"第 {row_number} 行缺少 GIGA Item Code（必填）")

        process_mode = field_value(row, "process_mode")
        validate_process_mode(process_mode, row_number)

        main_image_only = normalize_yes_no(field_value(row, "main_image_only"))
        if main_image_only:
            raise ValueError(f"第 {row_number} 行设置了只处理主图，但当前版本仅支持处理素材包内全部图片")

        actual_item_code = field_value(row, "actual_item_code") or expected_item_code
        output_prefix = field_value(row, "output_prefix") or field_value(row, "amazon_sku") or expected_item_code
        note_parts = []
        note = field_value(row, "note")
        if note:
            note_parts.append(note)
        if actual_item_code != expected_item_code:
            note_parts.append(f"模板 Item Code 为 {expected_item_code}，实际 Item Code 为 {actual_item_code}；按实际素材包处理。")

        source_zip = match_material_zip(input_dir, actual_item_code)
        if source_zip is None:
            missing_archives.append(
                f"第 {row_number} 行：期望 Item Code={expected_item_code}，实际 Item Code={actual_item_code}；"
                f"匹配规则=*{actual_item_code}*image+file*.zip；目录={input_dir}"
            )
            continue

        rows.append(
            {
                "expected_item_code": expected_item_code,
                "actual_item_code": actual_item_code,
                "output_prefix": output_prefix,
                "source_zip": source_zip,
                "note": " ".join(note_parts),
                "amazon_sku": field_value(row, "amazon_sku"),
                "product_url": field_value(row, "product_url"),
                "product_name": field_value(row, "product_name"),
            }
        )

    if missing_archives:
        raise RuntimeError("以下 Excel 行找不到素材 ZIP：\n" + "\n".join(missing_archives))
    return rows


def crop_cover(image: Image.Image, size: int) -> Image.Image:
    image = ImageOps.exif_transpose(image).convert("RGB")
    width, height = image.size
    scale = max(size / width, size / height)
    resized = image.resize((round(width * scale), round(height * scale)), Image.Resampling.LANCZOS)
    left = max((resized.width - size) // 2, 0)
    top = max((resized.height - size) // 2, 0)
    return resized.crop((left, top, left + size, top + size))


def iter_images(extract_dir: Path) -> Iterable[Path]:
    for path in sorted(extract_dir.rglob("*")):
        if path.is_file() and path.suffix.lower() in IMAGE_SUFFIXES and not path.name.startswith("._"):
            yield path


def safe_name(value: str) -> str:
    value = value.strip() or "unknown"
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", value)


def output_image_name(image_path: Path, used_names: set[str]) -> str:
    stem = image_path.stem.strip() or "image"
    name = f"{stem}.jpg"
    if name not in used_names:
        used_names.add(name)
        return name

    index = 2
    while True:
        candidate = f"{stem}_{index}.jpg"
        if candidate not in used_names:
            used_names.add(candidate)
            return candidate
        index += 1


def process_one_image(
    *,
    image_path: Path,
    product_extract_dir: Path,
    product_output_dir: Path,
    output_name: str,
    size: int,
    quality: int,
    expected_item_code: str,
    actual_item_code: str,
    output_prefix: str,
    source_zip: Path,
    note: str,
) -> tuple[dict, Path | None]:
    try:
        with Image.open(image_path) as image:
            original_size = f"{image.width}x{image.height}"
            processed = crop_cover(image, size)
        out_path = product_output_dir / output_name
        processed.save(out_path, "JPEG", quality=quality, optimize=True, progressive=True)
        return (
            {
                "期望Item Code": expected_item_code,
                "实际Item Code": actual_item_code,
                "输出前缀": output_prefix,
                "源压缩包": str(source_zip),
                "源图片": str(image_path.relative_to(product_extract_dir)),
                "输出图片": str(out_path),
                "原始尺寸": original_size,
                "输出尺寸": f"{size}x{size}",
                "处理状态": "success",
                "错误信息": "",
                "备注": note,
            },
            out_path,
        )
    except (UnidentifiedImageError, OSError) as exc:
        return (
            {
                "期望Item Code": expected_item_code,
                "实际Item Code": actual_item_code,
                "输出前缀": output_prefix,
                "源压缩包": str(source_zip),
                "源图片": str(image_path.relative_to(product_extract_dir)),
                "输出图片": "",
                "原始尺寸": "",
                "输出尺寸": "",
                "处理状态": "failed",
                "错误信息": str(exc),
                "备注": note,
            },
            None,
        )


def process_material_archives(
    *,
    input_dir: Path,
    output_dir: Path,
    report_dir: Path,
    size: int,
    quality: int,
    archive_rows: list[dict],
    workers: int | None = None,
) -> ProcessResult:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    extract_root = output_dir / f"_extracted_{timestamp}"
    final_root = output_dir
    report_dir.mkdir(parents=True, exist_ok=True)
    output_dir.mkdir(parents=True, exist_ok=True)
    extract_root.mkdir(parents=True, exist_ok=True)

    report_csv = report_dir / f"giga_image_process_report_{timestamp}.csv"
    export_zip = output_dir / f"giga_images_{timestamp}.zip"
    report_rows: list[dict] = []
    generated_images: list[Path] = []
    total_images = 0

    workers = workers or min(8, (os.cpu_count() or 4))
    workers = max(1, workers)

    for row in archive_rows:
        source_zip = Path(row["source_zip"])
        expected_item_code = str(row.get("expected_item_code") or "")
        actual_item_code = str(row.get("actual_item_code") or expected_item_code)
        output_prefix = safe_name(str(row.get("output_prefix") or actual_item_code or source_zip.stem))
        note = str(row.get("note") or "")
        product_extract_dir = extract_root / output_prefix
        product_output_dir = final_root / output_prefix
        product_extract_dir.mkdir(parents=True, exist_ok=True)
        product_output_dir.mkdir(parents=True, exist_ok=True)

        status = "success"
        error = ""
        image_jobs: list[tuple[Path, str]] = []
        try:
            with zipfile.ZipFile(source_zip) as zf:
                zf.extractall(product_extract_dir)
            used_names: set[str] = set()
            image_jobs = [(image_path, output_image_name(image_path, used_names)) for image_path in iter_images(product_extract_dir)]
            if image_jobs:
                with ThreadPoolExecutor(max_workers=workers) as executor:
                    futures = [
                        executor.submit(
                            process_one_image,
                            image_path=image_path,
                            product_extract_dir=product_extract_dir,
                            product_output_dir=product_output_dir,
                            output_name=output_name,
                            size=size,
                            quality=quality,
                            expected_item_code=expected_item_code,
                            actual_item_code=actual_item_code,
                            output_prefix=output_prefix,
                            source_zip=source_zip,
                            note=note,
                        )
                        for image_path, output_name in image_jobs
                    ]
                    for future in futures:
                        report_row, out_path = future.result()
                        report_rows.append(report_row)
                        if out_path is not None:
                            total_images += 1
                            generated_images.append(out_path)
            else:
                status = "failed"
                error = "压缩包内未找到可处理图片"
        except Exception as exc:  # noqa: BLE001 - 报告中保留单个素材包失败原因，继续处理其他包
            status = "failed"
            error = str(exc)

        if not image_jobs:
            report_rows.append(
                {
                    "期望Item Code": expected_item_code,
                    "实际Item Code": actual_item_code,
                    "输出前缀": output_prefix,
                    "源压缩包": str(source_zip),
                    "源图片": "",
                    "输出图片": "",
                    "原始尺寸": "",
                    "输出尺寸": "",
                    "处理状态": status,
                    "错误信息": error,
                    "备注": note,
                }
            )

    fieldnames = [
        "期望Item Code",
        "实际Item Code",
        "输出前缀",
        "源压缩包",
        "源图片",
        "输出图片",
        "原始尺寸",
        "输出尺寸",
        "处理状态",
        "错误信息",
        "备注",
    ]
    with report_csv.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(report_rows)

    with zipfile.ZipFile(export_zip, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for path in sorted(generated_images):
            zf.write(path, path.relative_to(final_root))

    return ProcessResult(export_zip=export_zip, report_csv=report_csv, image_count=total_images)


def process_from_template(
    *,
    template_path: Path,
    input_dir: Path,
    output_dir: Path,
    report_dir: Path,
    size: int,
    quality: int,
    sheet_name: str = DEFAULT_SHEET,
    workers: int | None = None,
) -> ProcessResult:
    rows = build_rows_from_template(template_path=template_path, input_dir=input_dir, sheet_name=sheet_name)
    return process_material_archives(
        input_dir=input_dir,
        output_dir=output_dir,
        report_dir=report_dir,
        size=size,
        quality=quality,
        archive_rows=rows,
        workers=workers,
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="处理 GIGA 素材包图片")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT_DIR)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--reports", type=Path, default=DEFAULT_REPORT_DIR)
    parser.add_argument("--size", type=int, default=1600)
    parser.add_argument("--quality", type=int, default=85)
    parser.add_argument("--workers", type=int, default=None, help="并发处理图片数量，默认自动按 CPU 选择，设为 1 可单线程处理")
    parser.add_argument("--dry-run", action="store_true", help="只检查 Excel 行和素材 ZIP 匹配，不处理图片")
    args = parser.parse_args()

    rows = build_rows_from_template(template_path=args.template, input_dir=args.input, sheet_name=args.sheet)
    print(f"template={args.template}")
    print(f"sheet={args.sheet}")
    print(f"matched_archives={len(rows)}")
    for row in rows:
        print(f"- {row['expected_item_code']} actual={row['actual_item_code']} -> {row['source_zip']}")

    if args.dry_run:
        return

    result = process_material_archives(
        input_dir=args.input,
        output_dir=args.output,
        report_dir=args.reports,
        size=args.size,
        quality=args.quality,
        archive_rows=rows,
        workers=args.workers,
    )
    print(f"export_zip={result.export_zip}")
    print(f"report_csv={result.report_csv}")
    print(f"image_count={result.image_count}")


if __name__ == "__main__":
    main()
